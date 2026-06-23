"""Shared file-to-text extraction used by both email attachment scanning and the file-upload pipeline."""

import io

MAX_CHARS = 5000


def extract_text_from_bytes(raw_bytes: bytes, ext: str, filename: str) -> str:
    """Extract plain text from file bytes. Returns up to MAX_CHARS chars, or an error marker string."""
    ext = ext.lower().lstrip('.')

    try:
        if ext in ('txt', 'md', 'csv'):
            return raw_bytes.decode('utf-8', errors='replace')[:MAX_CHARS]

        elif ext == 'pdf':
            try:
                import pdfplumber
                with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
                    return '\n'.join(p.extract_text() or '' for p in pdf.pages[:10])[:MAX_CHARS]
            except ImportError:
                return f"[PDF: {filename} — pdfplumber not installed]"

        elif ext in ('doc', 'docx'):
            try:
                import docx as python_docx
                doc = python_docx.Document(io.BytesIO(raw_bytes))
                return '\n'.join(p.text for p in doc.paragraphs if p.text)[:MAX_CHARS]
            except ImportError:
                return f"[Word: {filename} — python-docx not installed]"

        elif ext == 'xlsx':
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
                rows = []
                for sheet in list(wb.worksheets)[:3]:
                    rows.append(f"[Sheet: {sheet.title}]")
                    for row in sheet.iter_rows(max_row=100, values_only=True):
                        cells = [str(c) for c in row if c is not None]
                        if cells:
                            rows.append('\t'.join(cells))
                return '\n'.join(rows)[:MAX_CHARS]
            except ImportError:
                return f"[Excel: {filename} — openpyxl not installed]"

        elif ext == 'pptx':
            try:
                from pptx import Presentation
                prs = Presentation(io.BytesIO(raw_bytes))
                parts = []
                for i, slide in enumerate(prs.slides[:30]):
                    slide_texts = [s.text for s in slide.shapes if hasattr(s, 'text') and s.text.strip()]
                    if slide_texts:
                        parts.append(f"[Slide {i+1}] " + ' | '.join(slide_texts))
                return '\n'.join(parts)[:MAX_CHARS]
            except ImportError:
                return f"[PowerPoint: {filename} — python-pptx not installed]"

    except Exception as e:
        return f"[Error reading {filename}: {e}]"

    return ''
